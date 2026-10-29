"""
WB-KDIs training dataset: 8 invented SOEs x 10 fiscal years (2015-2024).

Built so that every chart and panel of the tool has data to draw:

* all required columns, plus every optional column of the standard schema
  (cash, receivables, debt split, FX and guaranteed debt, fuel and labour
  cost, COGS and other operating expenses for cost recovery, cash-flow lines,
  government-support flows and stocks, government arrears, VAT credits,
  exposure shares and provenance);
* equity, EBITDA and EBIT stay positive in every SOE-year, so no ratio is
  suppressed (ROE, debt/equity, debt/EBITDA, depreciation/EBITDA and the
  effective tax rate are defined everywhere);
* every SOE reports tariff/PSO compensation and at least one other support
  flow, so the support charts draw for whichever SOE is selected;
* the portfolio spans the three Z''-EM zones and moves between them over the
  decade (two SOEs slide from the grey zone into distress).

The statements are internally consistent: Total Assets = Total Liabilities +
Equity; EBITDA = EBIT + Depreciation; Operating Expense (incl. depreciation)
= Revenues - EBIT; Total Expense = Operating Expense + Interest + Tax; Net
Income = Revenues - Total Expense; COGS + Other Operating Expenses =
Operating Expense - Depreciation; Government Grants (booked in revenue) =
Tariff/PSO Compensation + Operating Subsidy; retained earnings and equity
roll forward with net income, dividends, equity injections and investment
grants (recorded in equity, as under SYSCOHADA).

Every figure is invented for training. These are not real enterprises.

    training_portfolio()      the Data sheet as a DataFrame
    training_extras()         Grants_detail, Source_lines, Validation, Definitions
    training_csv_text()       the CSV download (Data sheet only)
    training_excel_bytes()    the Excel download (Readme, Data and the extra sheets)
    write_files(folder)       writes both files (python training_data.py)
"""

from __future__ import annotations

import io
from pathlib import Path

import numpy as np
import pandas as pd

import schema

TRAINING_LABEL = "Training data: invented SOEs"
TRAINING_SHORT = "8 SOEs · 7 sectors · 2015–2024"
CSV_NAME = "WB-KDIs_SOE_training_data_2015-2024.csv"
XLSX_NAME = "WB-KDIs_SOE_training_data_2015-2024.xlsx"
YEARS = list(range(2015, 2025))
N = len(YEARS)
BN = 1e9
CURRENCY = "FCFA"
UNITS = "Absolute"


def _ramp(a, b):
    return list(np.linspace(a, b, N))


def _v(x, t):
    return x[t] if isinstance(x, (list, tuple, np.ndarray)) else x


# ---------------------------------------------------------------------------
# One dict per SOE. Lists run 2015 ... 2024; scalars hold for every year.
# rev, inv_grant, eq_inj in FCFA billions; everything else a ratio.
# ---------------------------------------------------------------------------
SOES = [
    dict(name="SOE A", sector="Energy", activity="Petroleum products importer and distributor",
         rev=[30.0, 31.6, 33.1, 35.4, 36.2, 31.0, 34.8, 44.6, 41.9, 42.8],
         ebitda_m=[0.150, 0.145, 0.135, 0.120, 0.105, 0.085, 0.080, 0.062, 0.066, 0.064],
         dep_m=0.030,
         tariff=[0.070, 0.075, 0.080, 0.090, 0.095, 0.120, 0.125, 0.210, 0.175, 0.170],
         opsub=[0.010, 0.010, 0.010, 0.012, 0.012, 0.020, 0.015, 0.015, 0.012, 0.012],
         inv_grant=[0.0, 0.0, 0.3, 0.0, 0.0, 0.0, 0.4, 0.0, 0.0, 0.0],
         eq_inj=[0.0] * N,
         turnover=[0.62, 0.62, 0.61, 0.61, 0.60, 0.52, 0.56, 0.68, 0.63, 0.62],
         ca=[0.48, 0.48, 0.49, 0.49, 0.49, 0.48, 0.49, 0.51, 0.50, 0.50],
         cr=[1.34, 1.30, 1.24, 1.17, 1.10, 1.00, 0.95, 0.87, 0.84, 0.82],
         payout=0.0, rate=0.072, st_cl=0.45, lt_ncl=0.80, fx_debt=0.70, float_debt=0.45,
         gg_lt=0.40, onlend_lt=0.10, fuel_opex=0.58, fx_cost=0.62, fx_rev=0.04,
         cogs_cash=0.86, labor_other=0.45,
         cash_ca=[0.14, 0.13, 0.12, 0.11, 0.10, 0.08, 0.08, 0.06, 0.06, 0.05], rec_ca=0.42, inv_ca=0.28, vat_ca=0.06,
         govarr_rec=_ramp(0.20, 0.45), e0=0.26, re0=0.05, tax_support=0.004),

    dict(name="SOE B", sector="Transport", activity="National railway (passenger and freight)",
         rev=[5.0, 5.2, 5.5, 5.8, 6.0, 5.1, 5.9, 6.5, 6.9, 7.2],
         ebitda_m=[0.240, 0.235, 0.232, 0.228, 0.226, 0.210, 0.220, 0.222, 0.225, 0.226],
         dep_m=0.165,
         tariff=[0.360, 0.370, 0.380, 0.390, 0.400, 0.460, 0.440, 0.450, 0.455, 0.460],
         opsub=[0.040, 0.040, 0.045, 0.045, 0.050, 0.070, 0.060, 0.060, 0.060, 0.060],
         inv_grant=[1.5, 1.2, 2.0, 1.8, 1.5, 1.0, 2.2, 2.5, 2.0, 1.8],
         eq_inj=[0.0, 0.0, 0.0, 0.0, 2.0, 0.0, 0.0, 0.0, 2.5, 0.0],
         turnover=[0.090, 0.089, 0.088, 0.087, 0.086, 0.073, 0.081, 0.084, 0.083, 0.082],
         ca=0.10, cr=[1.45, 1.42, 1.40, 1.38, 1.35, 1.22, 1.25, 1.24, 1.22, 1.20],
         payout=0.0, rate=0.032, st_cl=0.25, lt_ncl=0.85, fx_debt=0.55, float_debt=0.20,
         gg_lt=0.75, onlend_lt=0.40, fuel_opex=0.22, fx_cost=0.25, fx_rev=0.03,
         cogs_cash=0.55, labor_other=0.60,
         cash_ca=0.18, rec_ca=0.50, inv_ca=0.20, vat_ca=0.05,
         govarr_rec=_ramp(0.30, 0.55), e0=0.58, re0=-0.03, tax_support=0.010),

    dict(name="SOE C", sector="Power / Utilities", activity="Integrated electricity utility",
         rev=[150.0, 158.0, 168.0, 180.0, 192.0, 196.0, 212.0, 238.0, 250.0, 262.0],
         ebitda_m=[0.210, 0.205, 0.195, 0.180, 0.165, 0.150, 0.130, 0.100, 0.098, 0.100],
         dep_m=0.072,
         tariff=[0.080, 0.085, 0.095, 0.110, 0.125, 0.145, 0.165, 0.230, 0.220, 0.215],
         opsub=[0.010] * N,
         inv_grant=[4.0, 3.5, 5.0, 4.5, 4.0, 3.0, 6.0, 6.5, 5.0, 4.5],
         eq_inj=[0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 30.0, 0.0],
         turnover=[0.46, 0.46, 0.46, 0.45, 0.45, 0.44, 0.45, 0.48, 0.47, 0.47],
         ca=[0.28, 0.28, 0.29, 0.29, 0.30, 0.30, 0.31, 0.33, 0.33, 0.34],
         cr=[1.28, 1.23, 1.15, 1.07, 0.99, 0.91, 0.83, 0.74, 0.72, 0.70],
         payout=0.0, rate=0.066, st_cl=0.35, lt_ncl=0.82, fx_debt=0.80, float_debt=0.55,
         gg_lt=0.65, onlend_lt=0.20,
         fuel_opex=[0.40, 0.40, 0.41, 0.42, 0.42, 0.40, 0.43, 0.49, 0.47, 0.46], fx_cost=0.45, fx_rev=0.02,
         cogs_cash=0.78, labor_other=0.40,
         cash_ca=[0.12, 0.11, 0.10, 0.09, 0.08, 0.08, 0.07, 0.05, 0.05, 0.05], rec_ca=0.55, inv_ca=0.15, vat_ca=0.08,
         govarr_rec=_ramp(0.18, 0.48), e0=0.27, re0=0.04, tax_support=0.006),

    dict(name="SOE D", sector="Telecoms", activity="Fixed and mobile telecoms operator",
         rev=[45.0, 47.8, 50.5, 53.2, 56.0, 58.4, 61.5, 64.3, 67.0, 70.0],
         ebitda_m=[0.390, 0.392, 0.395, 0.393, 0.390, 0.388, 0.392, 0.385, 0.380, 0.382],
         dep_m=0.140,
         tariff=[0.010, 0.010, 0.012, 0.012, 0.013, 0.015, 0.014, 0.014, 0.013, 0.013],
         opsub=[0.0] * N,
         inv_grant=[0.0, 0.8, 0.0, 1.0, 0.0, 1.2, 0.0, 0.9, 0.0, 1.1],
         eq_inj=[0.0] * N,
         turnover=[0.52, 0.52, 0.52, 0.52, 0.52, 0.51, 0.52, 0.52, 0.52, 0.52],
         ca=0.34, cr=[1.80, 1.82, 1.85, 1.86, 1.88, 1.90, 1.92, 1.90, 1.92, 1.95],
         payout=0.65, rate=0.055, st_cl=0.20, lt_ncl=0.75, fx_debt=0.35, float_debt=0.30,
         gg_lt=0.10, onlend_lt=0.0, fuel_opex=0.06, fx_cost=0.35, fx_rev=0.08,
         cogs_cash=0.45, labor_other=0.40,
         cash_ca=0.30, rec_ca=0.45, inv_ca=0.10, vat_ca=0.04,
         govarr_rec=0.10, e0=0.52, re0=0.22, tax_support=0.0),

    dict(name="SOE E", sector="Agriculture", activity="Cotton development company",
         rev=[20.0, 24.2, 19.1, 26.0, 22.4, 18.3, 27.1, 30.2, 23.0, 25.4],
         ebitda_m=[0.110, 0.135, 0.075, 0.140, 0.095, 0.070, 0.145, 0.150, 0.080, 0.105],
         dep_m=0.032,
         tariff=[0.030, 0.025, 0.045, 0.020, 0.035, 0.050, 0.020, 0.020, 0.045, 0.035],
         opsub=[0.070, 0.060, 0.090, 0.060, 0.080, 0.100, 0.060, 0.055, 0.090, 0.075],
         inv_grant=[0.5, 0.0, 0.6, 0.0, 0.4, 0.0, 0.7, 0.0, 0.5, 0.0],
         eq_inj=[0.0] * N,
         turnover=[0.58, 0.66, 0.53, 0.68, 0.60, 0.50, 0.70, 0.74, 0.60, 0.64],
         ca=0.40, cr=[1.18, 1.25, 1.06, 1.26, 1.13, 1.02, 1.22, 1.25, 1.07, 1.12],
         payout=0.0, rate=0.075, st_cl=0.55, lt_ncl=0.70, fx_debt=0.40, float_debt=0.60,
         gg_lt=0.30, onlend_lt=0.15, fuel_opex=0.12, fx_cost=0.20, fx_rev=0.62,
         cogs_cash=0.80, labor_other=0.50,
         cash_ca=0.10, rec_ca=0.35, inv_ca=0.40, vat_ca=0.05,
         govarr_rec=0.25, e0=0.34, re0=0.03, tax_support=0.012),

    dict(name="SOE F", sector="Industry", activity="Cement producer",
         rev=[30.0, 31.2, 32.0, 32.6, 33.0, 30.5, 32.0, 32.4, 31.8, 31.5],
         ebitda_m=[0.240, 0.232, 0.222, 0.210, 0.198, 0.160, 0.165, 0.130, 0.120, 0.112],
         dep_m=0.065,
         tariff=[0.010, 0.010, 0.010, 0.012, 0.012, 0.015, 0.015, 0.020, 0.020, 0.020],
         opsub=[0.0, 0.0, 0.0, 0.0, 0.0, 0.010, 0.005, 0.0, 0.0, 0.0],
         inv_grant=[0.0] * N,
         eq_inj=[0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 1.5],
         turnover=[0.58, 0.58, 0.57, 0.56, 0.55, 0.50, 0.52, 0.52, 0.50, 0.49],
         ca=0.33, cr=[1.80, 1.74, 1.66, 1.58, 1.50, 1.30, 1.30, 1.18, 1.12, 1.08],
         payout=0.50, rate=0.068, st_cl=0.30, lt_ncl=0.85, fx_debt=0.60, float_debt=0.70,
         gg_lt=0.20, onlend_lt=0.0, fuel_opex=0.25, fx_cost=0.30, fx_rev=0.20,
         cogs_cash=0.70, labor_other=0.45,
         cash_ca=0.20, rec_ca=0.35, inv_ca=0.35, vat_ca=0.04,
         govarr_rec=0.08, e0=0.45, re0=0.15, tax_support=0.0),

    dict(name="SOE G", sector="Transport", activity="Urban bus and ferry company",
         rev=[25.0, 26.3, 27.5, 28.6, 29.8, 22.5, 26.0, 31.5, 33.8, 35.2],
         ebitda_m=[0.090, 0.088, 0.085, 0.080, 0.078, 0.062, 0.070, 0.064, 0.066, 0.065],
         dep_m=0.035,
         tariff=[0.100, 0.105, 0.110, 0.115, 0.120, 0.200, 0.160, 0.150, 0.150, 0.155],
         opsub=[0.050, 0.050, 0.050, 0.055, 0.055, 0.100, 0.070, 0.060, 0.060, 0.060],
         inv_grant=[1.0, 0.0, 1.5, 0.0, 1.0, 0.0, 2.0, 0.0, 1.5, 0.0],
         eq_inj=[0.0, 0.0, 0.0, 3.0, 0.0, 0.0, 4.0, 0.0, 0.0, 4.5],
         turnover=[0.50, 0.50, 0.49, 0.48, 0.48, 0.36, 0.41, 0.48, 0.50, 0.50],
         ca=0.22, cr=[0.95, 0.92, 0.88, 0.85, 0.82, 0.70, 0.72, 0.70, 0.68, 0.66],
         payout=0.0, rate=0.078, st_cl=0.40, lt_ncl=0.85, fx_debt=0.85, float_debt=0.50,
         gg_lt=0.85, onlend_lt=0.15, fuel_opex=0.35, fx_cost=0.45, fx_rev=0.10,
         cogs_cash=0.65, labor_other=0.55,
         cash_ca=[0.10, 0.09, 0.08, 0.08, 0.07, 0.05, 0.06, 0.05, 0.05, 0.04], rec_ca=0.40, inv_ca=0.20, vat_ca=0.07,
         govarr_rec=_ramp(0.35, 0.60), e0=0.24, re0=-0.06, tax_support=0.008),

    dict(name="SOE H", sector="Other", activity="Housing and real-estate company",
         rev=[9.0, 9.4, 9.8, 10.3, 10.7, 10.4, 11.0, 11.6, 12.1, 12.6],
         ebitda_m=[0.340, 0.345, 0.350, 0.352, 0.355, 0.330, 0.345, 0.350, 0.352, 0.355],
         dep_m=0.080,
         tariff=[0.025, 0.025, 0.026, 0.028, 0.028, 0.035, 0.030, 0.030, 0.030, 0.030],
         opsub=[0.0] * N,
         inv_grant=[0.4, 0.0, 0.0, 0.5, 0.0, 0.0, 0.6, 0.0, 0.0, 0.5],
         eq_inj=[0.0] * N,
         turnover=[0.36, 0.36, 0.36, 0.36, 0.36, 0.34, 0.35, 0.36, 0.36, 0.36],
         ca=0.42, cr=[2.20, 2.22, 2.25, 2.26, 2.28, 2.15, 2.22, 2.26, 2.30, 2.34],
         payout=0.60, rate=0.050, st_cl=0.15, lt_ncl=0.80, fx_debt=0.05, float_debt=0.25,
         gg_lt=0.10, onlend_lt=0.20, fuel_opex=0.03, fx_cost=0.05, fx_rev=0.0,
         cogs_cash=0.55, labor_other=0.50,
         cash_ca=0.28, rec_ca=0.30, inv_ca=0.35, vat_ca=0.03,
         govarr_rec=0.15, e0=0.55, re0=0.24, tax_support=0.0),
]

CIT_RATE = 0.30          # corporate income tax on profit after interest
MIN_TAX_OF_REVENUE = 0.005  # minimum (turnover) tax when there is no taxable profit
LT_REPAYMENT = 0.10      # share of opening long-term debt repaid each year (debt service)


def _soe_rows(p: dict) -> list:
    rows = []
    t0_ta = p["rev"][0] * BN / _v(p["turnover"], 0) / 1.04
    eq = p["e0"] * t0_ta
    re = p["re0"] * t0_ta
    cl0 = _v(p["ca"], 0) * t0_ta / _v(p["cr"], 0)
    tl0 = t0_ta - eq
    st0 = p["st_cl"] * cl0
    lt0 = p["lt_ncl"] * (tl0 - cl0)
    debt_prev, lt_prev = st0 + lt0, lt0
    fa_prev = t0_ta * (1 - _v(p["ca"], 0))
    nwc_prev = None
    for t, year in enumerate(YEARS):
        rev = p["rev"][t] * BN
        tariff = round(_v(p["tariff"], t) * rev)
        opsub = round(_v(p["opsub"], t) * rev)
        grants = tariff + opsub
        ebitda = _v(p["ebitda_m"], t) * rev
        dep = _v(p["dep_m"], t) * rev
        ebit = ebitda - dep
        opex = rev - ebit                      # includes depreciation
        interest = p["rate"] * debt_prev
        tax = max(CIT_RATE * max(0.0, ebit - interest), MIN_TAX_OF_REVENUE * rev)
        ni = ebit - interest - tax
        div = p["payout"] * ni if ni > 0 else 0.0
        inj = p["eq_inj"][t] * BN
        inv = p["inv_grant"][t] * BN
        re = re + ni - div
        eq = eq + ni - div + inj + inv         # investment grants recorded in equity
        ta = rev / _v(p["turnover"], t)
        tl = ta - eq
        ca = _v(p["ca"], t) * ta
        cl = ca / _v(p["cr"], t)
        st_debt = p["st_cl"] * cl
        lt_debt = p["lt_ncl"] * (tl - cl)
        debt = st_debt + lt_debt
        cash_ca, rec_ca, inv_ca, vat_ca = (_v(p[k], t) for k in ("cash_ca", "rec_ca", "inv_ca", "vat_ca"))
        cash, receivables, inventory, vat = cash_ca * ca, rec_ca * ca, inv_ca * ca, vat_ca * ca
        gov_arrears = _v(p["govarr_rec"], t) * receivables
        cash_costs = opex - dep
        cogs = p["cogs_cash"] * cash_costs
        other_opex = cash_costs - cogs
        fuel = _v(p["fuel_opex"], t) * opex
        labor = p["labor_other"] * other_opex
        fa = ta - ca
        capex = max(0.0, fa - fa_prev + dep)
        nwc = (ca - cash) - (cl - st_debt)
        ocf = ebitda - tax - (nwc - nwc_prev if nwc_prev is not None else 0.0)
        debt_service = interest + LT_REPAYMENT * lt_prev
        gg = p["gg_lt"] * lt_debt
        fx = p["fx_debt"] * debt
        rows.append({
            "Entity": p["name"], "Year": year, "Sector": p["sector"], "Currency": CURRENCY, "Units": UNITS,
            # income statement
            "Revenues": rev, "Operating Revenue": rev - grants, "Cost of Goods Sold": cogs,
            "Other Operating Expenses": other_opex, "Operating Expense": opex, "Fuel Cost": fuel, "Labor Cost": labor,
            "Depreciation": dep, "EBITDA": ebitda, "Operating Profits (EBIT)": ebit, "Interest Expense": interest,
            "Tax Expense": tax, "Total Expense": opex + interest + tax, "Net Income": ni,
            # balance sheet
            "Cash and Equivalents": cash, "Receivables": receivables, "Inventory": inventory,
            "Current Assets": ca, "Total Assets": ta, "Short Term Debt": st_debt, "Current Liabilities": cl,
            "Long Term Debt": lt_debt, "FX Debt": fx, "Government Guaranteed Debt": gg,
            "Total Liabilities": tl, "Equity": eq, "Retained Earnings": re,
            # cash flow
            "Operating Cash Flow": ocf, "Capex": capex, "Debt Service": debt_service,
            # government relationship (Government Grants = the support booked in revenue)
            "Government Grants": grants, "Tariff / PSO Compensation": tariff, "Operating Subsidy (Flow)": opsub,
            "Investment Grant (Flow)": inv, "Equity Injection (Flow)": inj,
            "Government Subsidies": opsub, "Capital Injection": inj,
            "On-lending": p["onlend_lt"] * lt_debt, "Concessional Loan (Stock)": p["onlend_lt"] * lt_debt,
            "Guarantee (Contingent)": gg, "Tax Support": _v(p["tax_support"], t) * rev,
            "Government Arrears Owed": gov_arrears, "VAT Credits": vat,
            # shock exposure shares (decimals)
            "Fuel Cost Share": fuel / opex, "FX Cost Share": p["fx_cost"], "FX Revenue Share": p["fx_rev"],
            "FX Debt Share": fx / tl, "Floating Rate Debt Share": p["float_debt"] * debt / tl,
        })
        debt_prev, lt_prev, fa_prev, nwc_prev = debt, lt_debt, fa, nwc
    return rows


_AMOUNT_ROUND = 0  # whole FCFA
_SHARE_COLS = ["Fuel Cost Share", "FX Cost Share", "FX Revenue Share", "FX Debt Share", "Floating Rate Debt Share"]


def _provenance(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    letter = df["Entity"].str[-1]
    df["Source File"] = "AR_" + df["Entity"].str.replace(" ", "_") + "_" + df["Year"].astype(str) + ".pdf"
    df["Source Page"] = ["pp. %d-%d" % (12 + (ord(c) - 65) * 2, 21 + (ord(c) - 65) * 2) for c in letter]
    df["Audit Status"] = np.select(
        [(df["Year"] == 2024) & df["Entity"].isin(["SOE A", "SOE G"]), (df["Year"] == 2024) & (df["Entity"] == "SOE C")],
        ["Unaudited", "Draft"], "Audited")
    df["Extraction Method"] = np.where(df["Entity"].isin(["SOE A", "SOE B"]), "Manual",
                                       np.where(df["Entity"].isin(["SOE E", "SOE G"]), "Semantic", "Machine"))
    df["Data Quality Flag"] = np.select(
        [(df["Entity"] == "SOE E") & (df["Year"] >= 2023), (df["Entity"] == "SOE B") & (df["Year"] <= 2016)],
        ["Proxy", "User assumption"], "Observed")
    return df


# column order of the Data sheet / CSV: identifiers, income statement, balance sheet,
# cash flow, government relationship, exposure shares, provenance
COLUMNS = [
    "Entity", "Year", "Sector", "Currency", "Units",
    "Revenues", "Operating Revenue", "Cost of Goods Sold", "Other Operating Expenses", "Operating Expense",
    "Fuel Cost", "Labor Cost", "Depreciation", "EBITDA", "Operating Profits (EBIT)", "Interest Expense",
    "Tax Expense", "Total Expense", "Net Income",
    "Cash and Equivalents", "Receivables", "Inventory", "Current Assets", "Total Assets",
    "Short Term Debt", "Current Liabilities", "Long Term Debt", "FX Debt", "Government Guaranteed Debt",
    "Total Liabilities", "Equity", "Retained Earnings",
    "Operating Cash Flow", "Capex", "Debt Service",
    "Government Grants", "Tariff / PSO Compensation", "Operating Subsidy (Flow)", "Investment Grant (Flow)",
    "Equity Injection (Flow)", "Government Subsidies", "Capital Injection", "On-lending",
    "Concessional Loan (Stock)", "Guarantee (Contingent)", "Tax Support", "Government Arrears Owed", "VAT Credits",
    "Fuel Cost Share", "FX Cost Share", "FX Revenue Share", "FX Debt Share", "Floating Rate Debt Share",
    "Source File", "Source Page", "Audit Status", "Extraction Method", "Data Quality Flag",
]


def training_portfolio() -> pd.DataFrame:
    """The Data sheet: 8 SOEs x 10 years, every column of the standard schema except Country."""
    df = pd.DataFrame([r for p in SOES for r in _soe_rows(p)])
    amounts = [c for c in df.columns if c not in _SHARE_COLS + ["Entity", "Year", "Sector", "Currency", "Units"]]
    df[amounts] = df[amounts].round(_AMOUNT_ROUND)
    # re-derive the identities from the rounded amounts so they hold to the unit
    df["Operating Expense"] = df["Revenues"] - df["Operating Profits (EBIT)"]
    df["EBITDA"] = df["Operating Profits (EBIT)"] + df["Depreciation"]
    df["Other Operating Expenses"] = df["Operating Expense"] - df["Depreciation"] - df["Cost of Goods Sold"]
    df["Total Expense"] = df["Operating Expense"] + df["Interest Expense"] + df["Tax Expense"]
    df["Net Income"] = df["Revenues"] - df["Total Expense"]
    df["Total Liabilities"] = df["Total Assets"] - df["Equity"]
    df["Government Grants"] = df["Tariff / PSO Compensation"] + df["Operating Subsidy (Flow)"]
    df["Operating Revenue"] = df["Revenues"] - df["Government Grants"]
    df[_SHARE_COLS] = df[_SHARE_COLS].round(4)
    df["Fuel Cost Share"] = (df["Fuel Cost"] / df["Operating Expense"]).round(4)
    df = _provenance(df)
    for c in COLUMNS:
        if c not in ("Entity", "Sector", "Currency", "Units", "Source File", "Source Page", "Audit Status",
                     "Extraction Method", "Data Quality Flag") and c not in _SHARE_COLS:
            df[c] = df[c].astype("int64")
    return df[COLUMNS].reset_index(drop=True)


# ---------------------------------------------------------------------------
# supplementary sheets (read by the upload when the Excel file is used)
# ---------------------------------------------------------------------------
_FLOW_CODES = [("tariff_pso_compensation", "Tariff / PSO Compensation", "Flow (in revenue)"),
               ("operating_subsidy", "Operating Subsidy (Flow)", "Flow (in revenue)"),
               ("investment_grant", "Investment Grant (Flow)", "Flow (in equity)"),
               ("equity_injection", "Equity Injection (Flow)", "Flow (in equity)")]
_OTHER_CODES = [("concessional_loan", "Concessional Loan (Stock)", "Stock"),
                ("guarantee", "Guarantee (Contingent)", "Contingent"),
                ("tax_support", "Tax Support", "Estimate")]


def training_extras(df: pd.DataFrame | None = None) -> dict:
    """Grants_detail, Source_lines, Validation and Definitions sheets (DataFrames)."""
    df = training_portfolio() if df is None else df
    grants = []
    for _, r in df.iterrows():
        for code, col, basis in _FLOW_CODES + _OTHER_CODES:
            grants.append({"Code": code, "Year": int(r["Year"]), "Value (FCFA, Ones)": int(r[col]), "Basis": basis,
                           "Source": f"Annual report {int(r['Year'])}, note on State support", "Entity": r["Entity"]})
    grants = pd.DataFrame(grants)

    lines = [("Revenues", "Chiffre d'affaires et produits (XB)"), ("Government Grants", "Subventions d'exploitation (TL)"),
             ("Operating Profits (EBIT)", "Résultat d'exploitation (XE)"), ("Net Income", "Résultat net (XI)"),
             ("Total Assets", "Total actif (BZ)"), ("Equity", "Capitaux propres (CP)")]
    src = []
    for ent, g in df.groupby("Entity", sort=True):
        g = g.set_index("Year")
        for col, label in lines:
            row = {"Code": f"{ent} · {col}", "SYSCOHADA line": label}
            row.update({str(y): int(g.at[y, col]) for y in YEARS})
            src.append(row)
    source_lines = pd.DataFrame(src)

    checks = [("Balance sheet: assets = liabilities + equity", "Total Assets", lambda r: r["Total Assets"] - r["Total Liabilities"] - r["Equity"]),
              ("EBITDA = EBIT + depreciation", "EBITDA", lambda r: r["EBITDA"] - r["Operating Profits (EBIT)"] - r["Depreciation"]),
              ("Net income = revenue − total expense", "Net Income", lambda r: r["Net Income"] - (r["Revenues"] - r["Total Expense"])),
              ("Grants = tariff/PSO compensation + operating subsidy", "Government Grants",
               lambda r: r["Government Grants"] - r["Tariff / PSO Compensation"] - r["Operating Subsidy (Flow)"])]
    val = []
    for _, r in df.iterrows():
        for name, item, f in checks:
            d = int(f(r))
            val.append({"Check": f"{r['Entity']}: {name}", "Year": int(r["Year"]), "Item": item, "Difference": d,
                        "Status": "OK" if d == 0 else "Review"})
    validation = pd.DataFrame(val)

    definitions = pd.DataFrame([{"Column": v["column"], "Meaning": v["label"], "Statement": v["statement"],
                                 "Used by": v["used_by"]} for v in schema.VARIABLES if v["column"] in df.columns])
    return {"Grants_detail": grants, "Source_lines": source_lines, "Validation": validation, "Definitions": definitions}


README_LINES = [
    "WB-KDIs Training : SOE Fiscal Risk — training dataset",
    "8 invented state-owned enterprises (SOE A to SOE H) in 7 sectors, fiscal years 2015–2024, amounts in FCFA (absolute units).",
    "All figures are invented for training. They are not the accounts of any real enterprise.",
    "",
    "How to use it",
    "1. Open the tool (SOE-Fiscal-Risk-Final.html, or streamlit run Home.py).",
    "2. On the Home screen, drop this file on 'Load data' (or click 'Choose a file').",
    "3. Every tab now draws from this file: KPI Dashboard, Distress Indicator, Expected Fiscal Cost, Shock Scenarios, Early Warning, Report.",
    "",
    "Sheets",
    "Data — one row per SOE per year (the sheet the tool reads). Column names follow the tool's standard schema.",
    "Grants_detail — government support by type and year (flows, stocks and contingent items kept apart).",
    "Source_lines — a few statement lines per SOE, as an audit trail.",
    "Validation — accounting identity checks (all OK).",
    "Definitions — what each column means.",
    "",
    "Conventions",
    "Operating Expense includes depreciation; EBIT = Revenues − Operating Expense; EBITDA = EBIT + Depreciation.",
    "Government Grants are booked in Revenues and equal Tariff / PSO Compensation + Operating Subsidy (Flow).",
    "Investment grants and equity injections go to equity, not revenue.",
    "Cost of Goods Sold + Other Operating Expenses = Operating Expense − Depreciation (cost recovery denominator).",
    "Exposure shares are decimals: 0.45 means 45%. FX Debt Share and Floating Rate Debt Share are shares of total liabilities.",
]


def training_csv_text(df: pd.DataFrame | None = None) -> str:
    df = training_portfolio() if df is None else df
    return df.to_csv(index=False, lineterminator="\n")


def training_excel_bytes(df: pd.DataFrame | None = None) -> bytes:
    """Readme, Data, Grants_detail, Source_lines, Validation and Definitions sheets."""
    from openpyxl.styles import Alignment, Font, PatternFill

    df = training_portfolio() if df is None else df
    extras = training_extras(df)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame({"Readme": README_LINES}).to_excel(xw, sheet_name="Readme", index=False)
        df.to_excel(xw, sheet_name="Data", index=False)
        for name, frame in extras.items():
            frame.to_excel(xw, sheet_name=name, index=False)
        head_fill = PatternFill("solid", fgColor="002244")
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head_fill
                cell.alignment = Alignment(vertical="center", wrap_text=ws.title == "Data")
            ws.freeze_panes = "C2" if ws.title == "Data" else "A2"
            for col in ws.columns:
                letter = col[0].column_letter
                head = str(col[0].value or "")
                if ws.title == "Readme":
                    ws.column_dimensions[letter].width = 120
                    continue
                width = max(len(head) if ws.title != "Data" else min(len(head), 16),
                            max((len(f"{c.value:,.0f}") if isinstance(c.value, (int, float)) and not isinstance(c.value, bool) and abs(c.value) >= 1000 else len(str(c.value or ""))) for c in col[1:]) if len(col) > 1 else 0)
                ws.column_dimensions[letter].width = min(48, max(9, width + 2))
                for c in col[1:]:
                    if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                        if head in _SHARE_COLS:
                            c.number_format = "0.00%"
                        elif head == "Year":
                            c.number_format = "0"
                        elif abs(c.value) >= 1000:
                            c.number_format = "#,##0"
            if ws.title == "Data":
                ws.row_dimensions[1].height = 45
        xw.book["Readme"]["A1"].font = Font(bold=True, color="FFFFFF", size=12)
    return _fixed_timestamps(buf.getvalue())


_FIXED_TIME = "2026-01-01T00:00:00Z"


def _fixed_timestamps(xlsx: bytes) -> bytes:
    """The same workbook bytes on every build (the page embeds the file and
    tests/check_snapshot_sync.py compares builds): fixed zip entry times and
    fixed created/modified dates in docProps/core.xml."""
    import re
    import zipfile

    src = zipfile.ZipFile(io.BytesIO(xlsx))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(<)", rb"\g<1>" + _FIXED_TIME.encode() + rb"\g<2>", data)
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            dst.writestr(zi, data)
    return out.getvalue()


def write_files(folder) -> tuple:
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    df = training_portfolio()
    (folder / CSV_NAME).write_text(training_csv_text(df), encoding="utf-8")
    (folder / XLSX_NAME).write_bytes(training_excel_bytes(df))
    return folder / CSV_NAME, folder / XLSX_NAME


if __name__ == "__main__":
    for path in write_files(Path(__file__).resolve().parent / "sample_data"):
        print("Wrote", path)
